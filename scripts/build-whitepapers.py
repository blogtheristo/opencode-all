#!/usr/bin/env python3
"""Build DWS whitepapers as colorful, branded A4 .docx — the single pipeline (replaces pandoc).

Implements DWS_DOCUMENT_STANDARD.md §22 (Whitepaper Standard):
  - A4 portrait, Aegis logo top-right, colored navy banner + two-tone title
  - Justified 12pt body; navy headings / blue sub-headings
  - Colored callout boxes (highlight=navy, disclaimer=amber, CTA=navy),
    KPI cards, tier cards, navy table headers — parsed from the source markup
  - Per-page footer: left=classification · center=copyright · right=big page number
  - Output filenames carry a numeric VERSION_TAG prefix (NN_WPn_...) so the
    latest build sorts together / is easy to find on Google Drive.

Sources: Learn/*.html (rich styled) + docs/lifetime-fleet/*.md.
Output:  build/whitepapers/ (git-ignored). Drag a .docx into Google Drive to get
         a Google Doc with the colorful formatting preserved — no pandoc needed.

Usage:    python3 scripts/build-whitepapers.py   |   npm run build:whitepapers
          VERSION_TAG=03 python3 scripts/build-whitepapers.py
Requires: python-docx  (pip install python-docx)
"""
import re, html, os
from html.parser import HTMLParser
from docx import Document
from docx.shared import Mm, Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO = os.path.join(REPO, "scripts/assets/aegis-logo.jpg")
OUT = os.path.join(REPO, "build/whitepapers")
VERSION_TAG = os.environ.get("VERSION_TAG", "02")
COPYRIGHT = "© 2026 Lifetime Oy · Y-tunnus 0772407-9 · All rights reserved."

# DWS palette
INK   = RGBColor(0x0F,0x2A,0x4F)   # deep navy (headings, titles)
ACC   = RGBColor(0x25,0x63,0xEB)   # accent blue (sub-headings, prices, KPI values)
GREY  = RGBColor(0x60,0x60,0x60)
WHITE = RGBColor(0xFF,0xFF,0xFF)
LTBLUE= RGBColor(0x93,0xC5,0xFD)   # light blue text on navy
NAVY_BG="0B1220"; AMBER_BG="FEF3C7"; BLUE_BG="EFF4FF"; CARD_BG="F8FAFC"; THEAD="0F2A4F"
AMBER_TX=RGBColor(0x7C,0x5E,0x10); AMBER_TI=RGBColor(0x92,0x40,0x0E)
BODY_ON_NAVY=RGBColor(0xE2,0xE8,0xF0)

# ---------------- HTML → lightweight DOM tree ----------------
VOID={'br','img','hr','meta','link','input','source'}
class El:
    __slots__=('tag','cls','children','parent')
    def __init__(self,tag,cls,parent): self.tag=tag; self.cls=cls; self.children=[]; self.parent=parent
class Tree(HTMLParser):
    def __init__(self): super().__init__(); self.root=El('root','',None); self.cur=self.root; self.skip=0
    def handle_starttag(self,tag,attrs):
        if tag in ('script','style'): self.skip+=1; return
        if self.skip: return
        if tag in VOID:
            if tag=='br': self.cur.children.append(('t','\n'))
            return
        cls=dict(attrs).get('class','') or ''
        el=El(tag,cls,self.cur); self.cur.children.append(el); self.cur=el
    def handle_startendtag(self,tag,attrs):
        if tag=='br' and not self.skip: self.cur.children.append(('t','\n'))
    def handle_endtag(self,tag):
        if tag in ('script','style'): self.skip=max(0,self.skip-1); return
        if self.skip or tag in VOID: return
        n=self.cur
        while n is not None and n.tag!=tag: n=n.parent
        if n is not None and n.parent is not None: self.cur=n.parent
    def handle_data(self,data):
        if self.skip: return
        self.cur.children.append(('t',data))

def kids(el): return [c for c in el.children if not isinstance(c,tuple)]
def clean(s): return re.sub(r'\s+',' ', html.unescape(s).replace('\xa0',' ')).strip()
def text_of(el):
    if isinstance(el,tuple): return el[1]
    return ''.join(text_of(c) for c in el.children)
def by_class(el,k):
    for c in kids(el):
        if k in (c.cls or ''): return c
    return None
def text_by_class(el,k):
    n=by_class(el,k); return clean(text_of(n)) if n else ''
def collect(el,tag,acc):
    for c in kids(el):
        if c.tag==tag: acc.append(c)
        else: collect(c,tag,acc)
def find(el,tag):
    if el.tag==tag: return el
    for c in kids(el):
        r=find(c,tag)
        if r is not None: return r
    return None

def split_h1(h1):
    main=''; accent=''
    for c in h1.children:
        if isinstance(c,tuple): main+=c[1]
        elif c.tag=='span': accent+=text_of(c)
        else: main+=text_of(c)
    return clean(main), clean(accent)

def table_rows(tbl):
    trs=[]; collect(tbl,'tr',trs); rows=[]
    for tr in trs:
        cells=[c for c in kids(tr) if c.tag in ('td','th')]
        rows.append([clean(text_of(c)) for c in cells])
    return [r for r in rows if any(c for c in r)]

# block model: ('title2',m,a)|('subtitle',t)|('h',lvl,t)|('p',t)|('li',t,ord)|('hr',)
#   |('table',rows)|('box',kind,title,paras)|('kpis',[(v,l)])|('tiers',[(n,price,desc,bullets)])
def parse_html(path):
    with open(path,encoding='utf-8') as f: src=f.read()
    t=Tree(); t.feed(src); body=find(t.root,'body') or t.root
    out=[]; emit(body,out); return out

def emit(el,out):
    for c in kids(el):
        cls=c.cls or ''; tag=c.tag
        if 'version-banner' in cls: continue            # replaced by standardized banner
        if tag=='h1': out.append(('title2',)+split_h1(c))
        elif tag in ('h2','h3','h4'): out.append(('h',int(tag[1]),clean(text_of(c))))
        elif tag=='p':
            if 'subtitle' in cls: out.append(('subtitle',clean(text_of(c))))
            else:
                tx=clean(text_of(c))
                if tx: out.append(('p',tx))
        elif tag in ('ul','ol'):
            for li in kids(c):
                if li.tag=='li': out.append(('li',clean(text_of(li)),tag=='ol'))
        elif tag=='table': out.append(('table',table_rows(c)))
        elif tag=='div':
            if 'disclaimer' in cls or 'footnote' in cls:
                out.append(('box','amber',None,[clean(text_of(c))]))
            elif 'highlight-box' in cls or 'cta-box' in cls:
                ti=''
                for h in kids(c):
                    if h.tag in ('h2','h3','h4'): ti=clean(text_of(h)); break
                paras=[clean(text_of(p)) for p in kids(c) if p.tag=='p' and clean(text_of(p))]
                out.append(('box','navy',ti or None,paras))
            elif 'kpi-row' in cls:
                items=[]
                for card in kids(c):
                    if 'kpi-card' in (card.cls or ''):
                        items.append((text_by_class(card,'value'), text_by_class(card,'label')))
                if items: out.append(('kpis',items))
            elif 'tier-grid' in cls:
                tiers=[]
                for card in kids(c):
                    if 'tier-card' in (card.cls or ''): tiers.append(parse_tier(card))
                if tiers: out.append(('tiers',tiers))
            else:
                emit(c,out)
        else:
            if kids(c): emit(c,out)

def parse_tier(card):
    name=''
    for h in kids(card):
        if h.tag in ('h3','h4'): name=clean(text_of(h)); break
    price=text_by_class(card,'price')
    ps=[clean(text_of(p)) for p in kids(card) if p.tag=='p' and clean(text_of(p))]
    desc=ps[0] if ps else ''
    bullets=[]
    for ul in kids(card):
        if ul.tag in ('ul','ol'):
            for li in kids(ul):
                if li.tag=='li': bullets.append(clean(text_of(li)))
    return (name,price,desc,bullets)

# ---------------- Markdown → blocks ----------------
def md_inline(s):
    s=re.sub(r'\*\*([^*]+)\*\*',r'\1',s); s=re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)',r'\1',s)
    s=re.sub(r'`([^`]+)`',r'\1',s); s=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',s); return s.strip()
def parse_markdown(path):
    with open(path,encoding='utf-8') as f: lines=f.read().split('\n')
    blocks=[]; i=0
    while i<len(lines):
        ln=lines[i].rstrip()
        if not ln.strip(): i+=1; continue
        if re.match(r'^#{1,4}\s',ln):
            lvl=len(ln)-len(ln.lstrip('#')); blocks.append(('h',min(lvl,4),md_inline(ln.lstrip('# ').strip()))); i+=1; continue
        if re.match(r'^(---+|\*\*\*+)\s*$',ln): blocks.append(('hr',)); i+=1; continue
        if ln.lstrip().startswith('|') and '|' in ln:
            tbl=[]
            while i<len(lines) and lines[i].lstrip().startswith('|'):
                row=lines[i].strip().strip('|')
                if re.match(r'^[\s:|-]+$',row): i+=1; continue
                tbl.append([md_inline(c.strip()) for c in row.split('|')]); i+=1
            if tbl: blocks.append(('table',tbl)); continue
        m=re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)',ln)
        if m: blocks.append(('li',md_inline(m.group(3)),bool(re.match(r'\d+\.',m.group(2))))); i+=1; continue
        para=[ln]; i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#{1,4}\s|\s*[-*]\s|\d+\.\s|\|)',lines[i]) and not re.match(r'^(---+|\*\*\*+)\s*$',lines[i]):
            para.append(lines[i].rstrip()); i+=1
        blocks.append(('p',md_inline(' '.join(para))))
    # promote first heading to two-tone title, the following sub-heading to subtitle
    if blocks and blocks[0][0]=='h': blocks[0]=('title2',blocks[0][2],'')
    if len(blocks)>1 and blocks[1][0]=='h' and blocks[1][1]>=3: blocks[1]=('subtitle',blocks[1][2])
    return blocks

# ---------------- docx helpers ----------------
def set_fill(cell,hexfill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hexfill); tcPr.append(shd)
def pad_cell(cell,pad=120):
    tcPr=cell._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
    for s in ('top','left','bottom','right'):
        e=OxmlElement('w:'+s); e.set(qn('w:w'),str(pad)); e.set(qn('w:type'),'dxa'); mar.append(e)
    tcPr.append(mar)
def add_page_field(par,size_pt):
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE')
    r=OxmlElement('w:r'); rpr=OxmlElement('w:rPr')
    sz=OxmlElement('w:sz'); sz.set(qn('w:val'),str(int(size_pt*2))); rpr.append(sz)
    col=OxmlElement('w:color'); col.set(qn('w:val'),'404040'); rpr.append(col)
    rpr.append(OxmlElement('w:b')); r.append(rpr)
    t=OxmlElement('w:t'); t.text='1'; r.append(t); fld.append(r); par._p.append(fld)

def style_doc(doc,classification):
    n=doc.styles['Normal']; n.font.name='Calibri'; n.font.size=Pt(12)
    n.paragraph_format.line_spacing=1.15; n.paragraph_format.space_after=Pt(6)
    s=doc.sections[0]; s.orientation=WD_ORIENT.PORTRAIT
    s.page_width=Mm(210); s.page_height=Mm(297)
    s.top_margin=Mm(20); s.bottom_margin=Mm(20); s.left_margin=Mm(22); s.right_margin=Mm(22)
    s.header_distance=Mm(12); s.footer_distance=Mm(10)
    hp=s.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    hp.add_run().add_picture(LOGO, width=Inches(1.15))
    fp=s.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.LEFT
    ts=fp.paragraph_format.tab_stops
    ts.add_tab_stop(Mm(83),WD_TAB_ALIGNMENT.CENTER); ts.add_tab_stop(Mm(166),WD_TAB_ALIGNMENT.RIGHT)
    rc=fp.add_run(classification); rc.font.size=Pt(7.5); rc.font.color.rgb=GREY; rc.bold=True
    fp.add_run('\t'); rcp=fp.add_run(COPYRIGHT); rcp.font.size=Pt(7.5); rcp.font.color.rgb=GREY
    fp.add_run('\t'); add_page_field(fp,13)

def add_banner(doc,title,version):
    t=doc.add_table(rows=1,cols=1); cell=t.cell(0,0); set_fill(cell,NAVY_BG); pad_cell(cell,150)
    p=cell.paragraphs[0]; r=p.add_run(title); r.bold=True; r.font.size=Pt(13); r.font.color.rgb=WHITE
    p2=cell.add_paragraph(); r2=p2.add_run(version); r2.font.size=Pt(9); r2.font.color.rgb=LTBLUE
    doc.add_paragraph()

def add_box(doc,kind,title,paras):
    fill=AMBER_BG if kind=='amber' else NAVY_BG
    ti_c=AMBER_TI if kind=='amber' else LTBLUE
    bd_c=AMBER_TX if kind=='amber' else BODY_ON_NAVY
    t=doc.add_table(rows=1,cols=1); cell=t.cell(0,0); set_fill(cell,fill); pad_cell(cell,140)
    p0=cell.paragraphs[0]; started=False
    if title:
        r=p0.add_run(title); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=ti_c; started=True
    for para in paras:
        p=p0 if not started else cell.add_paragraph(); started=True
        r=p.add_run(para); r.font.color.rgb=bd_c; r.font.size=Pt(11)
    doc.add_paragraph()

def add_kpis(doc,items):
    t=doc.add_table(rows=1,cols=len(items))
    for i,(val,lab) in enumerate(items):
        cell=t.cell(0,i); set_fill(cell,BLUE_BG); pad_cell(cell,140)
        p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(val); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=ACC
        p2=cell.add_paragraph(); p2.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r2=p2.add_run(lab); r2.font.size=Pt(8.5); r2.font.color.rgb=GREY
    doc.add_paragraph()

def add_tiers(doc,tiers):
    t=doc.add_table(rows=1,cols=len(tiers))
    for i,(name,price,desc,bullets) in enumerate(tiers):
        cell=t.cell(0,i); set_fill(cell,CARD_BG); pad_cell(cell,140)
        p=cell.paragraphs[0]; r=p.add_run(name); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=INK
        if price:
            pp=cell.add_paragraph(); rp=pp.add_run(price); rp.bold=True; rp.font.size=Pt(12); rp.font.color.rgb=ACC
        if desc:
            pd=cell.add_paragraph(); rd=pd.add_run(desc); rd.font.size=Pt(9); rd.font.color.rgb=GREY
        for b in bullets:
            pb=cell.add_paragraph(b,style='List Bullet')
            for r in pb.runs: r.font.size=Pt(9.5)
    doc.add_paragraph()

def add_title(doc,main,accent):
    p=doc.add_paragraph()
    r=p.add_run(main+(' ' if accent else '')); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=INK
    if accent:
        r2=p.add_run(accent); r2.bold=True; r2.font.size=Pt(24); r2.font.color.rgb=ACC

def add_data_table(doc,rows):
    ncol=max(len(r) for r in rows); t=doc.add_table(rows=0,cols=ncol); t.style='Light Grid Accent 1'
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for ci in range(ncol):
            cells[ci].text=row[ci] if ci<len(row) else ''
            if ri==0:
                set_fill(cells[ci],THEAD)
                for para in cells[ci].paragraphs:
                    for run in para.runs: run.bold=True; run.font.color.rgb=WHITE
    doc.add_paragraph()

def render(blocks,out,version,classification,title_display):
    doc=Document(); style_doc(doc,classification); add_banner(doc,title_display,version)
    for b in blocks:
        k=b[0]
        if k=='title2': add_title(doc,b[1],b[2])
        elif k=='subtitle':
            p=doc.add_paragraph(); r=p.add_run(b[1]); r.italic=True; r.font.size=Pt(12.5); r.font.color.rgb=GREY
        elif k=='h':
            h=doc.add_heading(b[2],level=min(b[1],4))
            col=INK if b[1]<=2 else ACC
            for r in h.runs: r.font.color.rgb=col
        elif k=='p':
            p=doc.add_paragraph(b[1]); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        elif k=='li':
            p=doc.add_paragraph(b[1],style='List Number' if b[2] else 'List Bullet'); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        elif k=='hr':
            p=doc.add_paragraph(); p.add_run('—'*30).font.color.rgb=RGBColor(0xC0,0xC0,0xC0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif k=='table': add_data_table(doc,b[1])
        elif k=='box': add_box(doc,b[1],b[2],b[3])
        elif k=='kpis': add_kpis(doc,b[1])
        elif k=='tiers': add_tiers(doc,b[1])
    doc.save(out); return out

# (source, kind, WPnum, slug, display-title, version-template{cls}, classification)
JOBS=[
  ("Learn/sovereign-industrial-intelligence-whitepaper-v1.0.html","html",1,"Sovereign_Industrial_Intelligence",
   "European Sovereign Industrial Intelligence",
   "Whitepaper · Version 1.0 · 2026 · Lifetime Oy · Classification: {cls}","Confidential"),
  ("Learn/KYA-IDV-CERT-Whitepaper-v1.0.html","html",2,"KYA_IDV_CERT",
   "KYA Standard — IDV Provider Certification",
   "Whitepaper · KYA-IDV-CERT · Version 1.0 · 2026 · Lifetime Oy · Classification: {cls}","Confidential"),
  ("Learn/2026-03-17-shufti-kya-whitepaper_v1.0.html","html",3,"Shufti_KYA",
   "KYA Standard — Shufti IDV Certification",
   "Whitepaper · Version 1.0 · 2026-03-17 · Lifetime Oy · Classification: {cls}","Confidential"),
  ("docs/lifetime-fleet/WHITEPAPER_COMPLIANT_BY_CONSTRUCTION.md","md",4,"Compliant_by_Construction",
   "Compliant by Construction",
   "Whitepaper · Version 1.0 · 2026-06-24 · Lifetime Oy / DWS IQ 6 · Classification: {cls}","Public"),
]

def main():
    os.makedirs(OUT,exist_ok=True)
    for src,kind,num,slug,title,vtmpl,classification in JOBS:
        path=os.path.join(REPO,src)
        blocks=parse_html(path) if kind=='html' else parse_markdown(path)
        version=vtmpl.format(cls=classification)
        name=f"{VERSION_TAG}_WP{num}_{slug}_A4.docx"
        render(blocks,os.path.join(OUT,name),version,classification,title)
        print(f"{name}: {len(blocks)} blocks, {classification}")

if __name__=="__main__":
    main()
